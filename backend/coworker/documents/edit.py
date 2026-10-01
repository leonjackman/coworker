"""In-place edits for docx / xlsx / pptx / pdf, operating on raw bytes.

Each function takes the existing ``bytes`` and a list of operation dicts, and
returns the new ``bytes``. Operations are applied in order; an unknown/invalid
operation raises :class:`DocumentError` and the caller leaves the file untouched
(the tool builds the new bytes first, then the workspace performs a single
atomic write).

Supported operations (``{"op": ...}``):

* docx — ``replace_text`` (old/new/count), ``append_paragraph`` (text/style),
  ``delete_paragraph`` (index).
* xlsx — ``set_cell`` (sheet/cell/value), ``append_row`` (sheet/values),
  ``delete_rows`` (sheet/start/count), ``add_sheet`` (name),
  ``delete_sheet`` (name).
* pptx — ``add_slide`` (title/bullets), ``delete_slide`` (index, 1-based),
  ``replace_text`` (old/new).
* pdf  — ``delete_pages`` (pages), ``rotate`` (pages/degrees),
  ``set_metadata`` (title/author), ``encrypt`` (password).
"""

from __future__ import annotations

import io
from typing import Any

from .common import DocumentError


def edit_document_bytes(fmt: str, data: bytes, operations: list[dict[str, Any]]) -> bytes:
    if not isinstance(operations, list) or not operations:
        raise DocumentError("operations must be a non-empty list", "bad_operations")
    if fmt == "docx":
        return _edit_docx(data, operations)
    if fmt == "xlsx":
        return _edit_xlsx(data, operations)
    if fmt == "pptx":
        return _edit_pptx(data, operations)
    if fmt == "pdf":
        return _edit_pdf(data, operations)
    raise DocumentError(f"Cannot edit format: {fmt}", "unsupported_format")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _edit_docx(data: bytes, operations: list[dict[str, Any]]) -> bytes:
    try:
        import docx
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"python-docx unavailable: {exc}", "dependency_missing") from exc
    document = docx.Document(io.BytesIO(data))
    for op in operations:
        kind = str(op.get("op") or "")
        if kind == "replace_text":
            old = str(op.get("old") or "")
            new = str(op.get("new") or "")
            if not old:
                raise DocumentError("replace_text requires 'old'", "bad_operations")
            _docx_replace(document, old, new, op.get("count"))
        elif kind == "append_paragraph":
            text = str(op.get("text") or "")
            style = op.get("style")
            document.add_paragraph(text, style=str(style)) if style else document.add_paragraph(text)
        elif kind == "delete_paragraph":
            index = int(op.get("index"))
            paragraphs = document.paragraphs
            if index < 0 or index >= len(paragraphs):
                raise DocumentError(f"delete_paragraph index {index} out of range", "bad_operations")
            element = paragraphs[index]._element
            element.getparent().remove(element)
        else:
            raise DocumentError(f"Unsupported docx operation: {kind}", "bad_operations")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_replace(document: Any, old: str, new: str, count: Any) -> None:
    limit = int(count) if count not in (None, "") else 0  # 0 = replace all
    replaced = 0

    def replace_in_paragraph(paragraph: Any) -> None:
        nonlocal replaced
        for run in paragraph.runs:
            if old not in run.text:
                continue
            if limit == 0:
                run.text = run.text.replace(old, new)
                replaced += 1
            else:
                while replaced < limit and old in run.text:
                    run.text = run.text.replace(old, new, 1)
                    replaced += 1

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def _edit_xlsx(data: bytes, operations: list[dict[str, Any]]) -> bytes:
    try:
        import openpyxl
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"openpyxl unavailable: {exc}", "dependency_missing") from exc
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    for op in operations:
        kind = str(op.get("op") or "")
        if kind == "set_cell":
            sheet = _sheet(workbook, op)
            sheet[str(op.get("cell") or "A1")] = op.get("value")
        elif kind == "append_row":
            sheet = _sheet(workbook, op)
            sheet.append(list(op.get("values") or []))
        elif kind == "delete_rows":
            sheet = _sheet(workbook, op)
            start = int(op.get("start") or 1)
            count = int(op.get("count") or 1)
            sheet.delete_rows(max(1, start), max(1, count))
        elif kind == "add_sheet":
            name = str(op.get("name") or "Sheet")
            if name in workbook.sheetnames:
                raise DocumentError(f"sheet already exists: {name}", "bad_operations")
            workbook.create_sheet(title=name[:31])
        elif kind == "delete_sheet":
            name = str(op.get("name") or "")
            if name not in workbook.sheetnames:
                raise DocumentError(f"sheet not found: {name}", "bad_operations")
            if len(workbook.sheetnames) <= 1:
                raise DocumentError("cannot delete the only sheet", "bad_operations")
            del workbook[name]
        else:
            raise DocumentError(f"Unsupported xlsx operation: {kind}", "bad_operations")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet(workbook: Any, op: dict[str, Any]) -> Any:
    name = str(op.get("sheet") or "")
    if name:
        if name not in workbook.sheetnames:
            raise DocumentError(f"sheet not found: {name}", "bad_operations")
        return workbook[name]
    return workbook.active


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _edit_pptx(data: bytes, operations: list[dict[str, Any]]) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"python-pptx unavailable: {exc}", "dependency_missing") from exc
    presentation = Presentation(io.BytesIO(data))
    for op in operations:
        kind = str(op.get("op") or "")
        if kind == "add_slide":
            blank = presentation.slide_layouts[6]
            slide = presentation.slides.add_slide(blank)
            title = str(op.get("title") or "")
            if title:
                box = slide.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(9), Inches(1))
                box.text_frame.text = title
            bullets = op.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            if bullets:
                box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(8.4), Inches(4.5))
                frame = box.text_frame
                frame.word_wrap = True
                for index, item in enumerate(bullets):
                    para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    para.text = f"• {item}"
        elif kind == "delete_slide":
            index = int(op.get("index") or 0)
            count = len(presentation.slides)
            if index < 1 or index > count:
                raise DocumentError(f"delete_slide index {index} out of range (1-{count})", "bad_operations")
            if count <= 1:
                raise DocumentError("cannot delete the only slide", "bad_operations")
            slide_id_list = presentation.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public API
            slides = list(slide_id_list)
            slide_id_list.remove(slides[index - 1])
        elif kind == "replace_text":
            old = str(op.get("old") or "")
            new = str(op.get("new") or "")
            if not old:
                raise DocumentError("replace_text requires 'old'", "bad_operations")
            _pptx_replace(presentation, old, new)
        else:
            raise DocumentError(f"Unsupported pptx operation: {kind}", "bad_operations")
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx_replace(presentation: Any, old: str, new: str) -> None:
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if old in run.text:
                                    run.text = run.text.replace(old, new)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _edit_pdf(data: bytes, operations: list[dict[str, Any]]) -> bytes:
    try:
        from pypdf import PdfReader, PdfWriter
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"pypdf unavailable: {exc}", "dependency_missing") from exc
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)

    for op in operations:
        kind = str(op.get("op") or "")
        if kind == "delete_pages":
            indices = _page_indices(op.get("pages"), len(writer.pages))
            for index in sorted(indices, reverse=True):
                writer.remove_page(index)
        elif kind == "rotate":
            degrees = int(op.get("degrees") or 90)
            for index in _page_indices(op.get("pages"), len(writer.pages)):
                writer.pages[index].rotate(degrees % 360)
        elif kind == "set_metadata":
            metadata = dict(reader.metadata or {})
            if op.get("title") is not None:
                metadata["/Title"] = str(op.get("title"))
            if op.get("author") is not None:
                metadata["/Author"] = str(op.get("author"))
            writer.add_metadata(metadata)
        elif kind == "encrypt":
            password = str(op.get("password") or "")
            if not password:
                raise DocumentError("encrypt requires 'password'", "bad_operations")
            writer.encrypt(user_password=password, owner_password=password or None)
        else:
            raise DocumentError(f"Unsupported pdf operation: {kind}", "bad_operations")

    if len(writer.pages) == 0:
        raise DocumentError("the operation would leave the PDF with no pages", "bad_operations")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _page_indices(spec: Any, page_count: int) -> list[int]:
    """Parse a 1-based page spec like ``"1,3-5"`` (or an int) into 0-based indices."""
    if spec is None or spec == "":
        return list(range(page_count))
    tokens: list[str] = []
    if isinstance(spec, int):
        tokens = [str(spec)]
    elif isinstance(spec, (list, tuple)):
        tokens = [str(item) for item in spec]
    else:
        tokens = [part.strip() for part in str(spec).split(",")]
    indices: list[int] = []
    for token in tokens:
        if not token:
            continue
        if "-" in token:
            start_s, _, end_s = token.partition("-")
            start = int(start_s)
            end = int(end_s) if end_s else page_count
            for number in range(start, end + 1):
                indices.append(number - 1)
        else:
            indices.append(int(token) - 1)
    valid = sorted({i for i in indices if 0 <= i < page_count})
    if not valid:
        raise DocumentError(f"page selection '{spec}' matched no pages", "bad_operations")
    return valid
