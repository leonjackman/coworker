"""Read (extract) content from docx / xlsx / pptx / pdf.

Every reader takes raw ``bytes`` so the same code path serves both a workspace
file (read then passed as bytes) and a chat attachment (a base64 data URL).
Each returns a bounded, JSON-serialisable dict with a readable ``text`` field
plus light structure and metadata.
"""

from __future__ import annotations

import io
from typing import Any

from .common import (
    DEFAULT_READ_MAX_CHARS,
    MAX_READ_MAX_CHARS,
    DocumentError,
    stringify_cell,
    truncate_text,
)

# Per-extractor hard bounds so a pathological document cannot blow up memory or
# the model context before the char cap is applied.
MAX_SHEET_ROWS = 500
MAX_SHEET_COLS = 60
MAX_SLIDES = 200
MAX_PDF_PAGES = 200


def read_document_bytes(data: bytes, fmt: str, *, max_chars: int = DEFAULT_READ_MAX_CHARS) -> dict[str, Any]:
    """Extract ``fmt`` content from ``data`` into a bounded dict."""
    max_chars = max(1, min(int(max_chars or DEFAULT_READ_MAX_CHARS), MAX_READ_MAX_CHARS))
    if fmt == "docx":
        return _read_docx(data, max_chars)
    if fmt == "xlsx":
        return _read_xlsx(data, max_chars)
    if fmt == "pptx":
        return _read_pptx(data, max_chars)
    if fmt == "pdf":
        return _read_pdf(data, max_chars)
    raise DocumentError(f"Unsupported document type: {fmt}", "unsupported_format")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _read_docx(data: bytes, max_chars: int) -> dict[str, Any]:
    try:
        import docx  # python-docx
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"python-docx unavailable: {exc}", "dependency_missing") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Not a valid .docx file: {exc}", "corrupt_document") from exc

    lines: list[str] = []
    paragraphs = 0
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        paragraphs += 1
        style = ""
        try:
            style = (para.style.name or "") if para.style is not None else ""
        except Exception:  # noqa: BLE001
            style = ""
        prefix = "# " if "Heading 1" in style else ("## " if style.startswith("Heading") else "")
        lines.append(f"{prefix}{text}")

    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([stringify_cell(cell.text.strip()) for cell in row.cells])
        tables.append(rows)
        lines.append("")
        for row in rows:
            lines.append("| " + " | ".join(row) + " |")

    text, truncated = truncate_text("\n".join(lines).strip(), max_chars)
    return {
        "format": "docx",
        "text": text,
        "metadata": {"paragraphs": paragraphs, "tables": len(tables)},
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def _read_xlsx(data: bytes, max_chars: int) -> dict[str, Any]:
    try:
        import openpyxl
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"openpyxl unavailable: {exc}", "dependency_missing") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Not a valid .xlsx file: {exc}", "corrupt_document") from exc

    lines: list[str] = []
    sheets_meta: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            row_count = 0
            for row in sheet.iter_rows(values_only=True, max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLS):
                rows = [stringify_cell(cell) for cell in row]
                while rows and rows[-1] == "":
                    rows.pop()
                if not rows:
                    continue
                row_count += 1
                lines.append("\t".join(rows))
            sheets_meta.append({"name": sheet.title, "rows": min(row_count, MAX_SHEET_ROWS)})
            lines.append("")
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass

    text, truncated = truncate_text("\n".join(lines).strip(), max_chars)
    return {
        "format": "xlsx",
        "text": text,
        "metadata": {"sheets": sheets_meta},
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _read_pptx(data: bytes, max_chars: int) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"python-pptx unavailable: {exc}", "dependency_missing") from exc
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Not a valid .pptx file: {exc}", "corrupt_document") from exc

    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            lines.append(f"... ({len(presentation.slides) - MAX_SLIDES} more slides truncated)")
            break
        lines.append(f"## Slide {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
        notes = _slide_notes(slide)
        if notes:
            lines.append(f"[notes] {notes}")
        lines.append("")

    text, truncated = truncate_text("\n".join(lines).strip(), max_chars)
    return {
        "format": "pptx",
        "text": text,
        "metadata": {"slides": len(presentation.slides)},
        "truncated": truncated,
    }


def _slide_notes(slide: Any) -> str:
    try:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            return slide.notes_slide.notes_text_frame.text.strip()
    except Exception:  # noqa: BLE001
        return ""
    return ""


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _read_pdf(data: bytes, max_chars: int) -> dict[str, Any]:
    try:
        import pdfplumber
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"pdfplumber unavailable: {exc}", "dependency_missing") from exc

    lines: list[str] = []
    metadata: dict[str, Any] = {}
    page_count = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            metadata = {
                "pages": page_count,
                "title": (pdf.metadata or {}).get("Title") or "",
                "author": (pdf.metadata or {}).get("Author") or "",
            }
            for index, page in enumerate(pdf.pages, start=1):
                if index > MAX_PDF_PAGES:
                    lines.append(f"... ({page_count - MAX_PDF_PAGES} more pages truncated)")
                    break
                lines.append(f"## Page {index}")
                page_text = page.extract_text() or ""
                lines.append(page_text.strip())
                lines.append("")
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Failed to read PDF: {exc}", "corrupt_document") from exc

    text, truncated = truncate_text("\n".join(lines).strip(), max_chars)
    return {
        "format": "pdf",
        "text": text,
        "metadata": metadata,
        "truncated": truncated,
    }


def extract_text(data: bytes, fmt: str, *, max_chars: int = DEFAULT_READ_MAX_CHARS) -> tuple[str, bool]:
    """Convenience: just the readable text (used for chat attachments)."""
    result = read_document_bytes(data, fmt, max_chars=max_chars)
    return str(result.get("text") or ""), bool(result.get("truncated"))
